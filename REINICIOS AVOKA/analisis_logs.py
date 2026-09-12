import os
import glob
import re
import pandas as pd
import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def parse_all_logs():
    """
    Lee todos los archivos .log de la carpeta y extrae los datos 
    identificando automáticamente el servidor de origen.
    """
    log_files = glob.glob("*.log")
    if not log_files:
        print("No se encontraron archivos .log en el directorio actual.")
        return pd.DataFrame()

    print(f"Archivos de log detectados: {len(log_files)}")
    
    data = []
    line_regex = re.compile(
        r'^(\d{2}:\d{2}:\d{2}\s+(?:AM|PM))\s+(\d+)\s+(\d+)\s+([\d\.]+)\s+([\d\.]+)\s+([\d\.]+)\s+([\d\.]+)\s+([\d\.]+)\s+(\d+)\s+(.+)$'
    )
    date_regex = re.compile(r'(\d{2}/\d{2}/\d{4})')

    for file_path in log_files:
        server_from_file = os.path.basename(file_path).split('_pidstat')[0] if '_pidstat' in file_path else "Servidor_Desconocido"
        current_server = server_from_file
        current_date = "Desconocida"

        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue

                if "Linux" in line:
                    match_date = date_regex.search(line)
                    if match_date:
                        current_date = match_date.group(1)
                    parts = line.split()
                    if len(parts) >= 3 and "(" in parts[2]:
                        current_server = parts[2].replace("(", "").replace(")", "")
                    continue

                match = line_regex.match(line)
                if match:
                    timestamp_str, uid, pid, usr, system, guest, wait, cpu, cpu_id, cmd = match.groups()
                    data.append({
                        'Server': current_server,
                        'Date': current_date,
                        'Timestamp': f"{current_date} {timestamp_str}",
                        'UID': int(uid),
                        'PID': int(pid),
                        'USR': float(usr),
                        'SYSTEM': float(system),
                        'WAIT': float(wait),
                        'CPU_PERCENT': float(cpu),
                        'CPU_ID': int(cpu_id),
                        'Command': cmd.strip()
                    })

    df = pd.DataFrame(data)
    return df

def generate_excel_for_server(server_name, df_server):
    """
    Genera el reporte de Excel estructurado para un servidor específico.
    """
    filename = f"Analisis_CPU_{server_name}.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Eliminar hoja predeterminada

    # Estilos
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=15, bold=True, color="1F4E79")
    sub_title_font = Font(name="Calibri", size=12, bold=True, color="333333")
    regular_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )

    # 1. RESUMEN EJECUTIVO
    ws1 = wb.create_sheet(title="Resumen Ejecutivo")
    ws1.views.sheetView[0].showGridLines = True
    
    ws1["A1"] = f"REPORTE DE RENDIMIENTO CPU - SERVIDOR: {server_name}"
    ws1["A1"].font = title_font
    
    ws1["A3"] = "1. Métricas Globales del Servidor"
    ws1["A3"].font = sub_title_font
    
    max_cpu_row = df_server.loc[df_server['CPU_PERCENT'].idxmax()] if not df_server.empty else None
    max_cmd = f"{max_cpu_row['Command']} (PID {max_cpu_row['PID']} - {max_cpu_row['CPU_PERCENT']}%)" if max_cpu_row is not None else "N/A"
    
    metrics = [
        ("Servidor Analizado", server_name),
        ("Total Registros Muestreados", len(df_server)),
        ("Procesos Diferentes Registrados", df_server['Command'].nunique()),
        ("Pico Máximo de Consumo de CPU", max_cmd),
        ("Promedio de CPU General (Todos los procesos)", f"{round(df_server['CPU_PERCENT'].mean(), 2)}%")
    ]
    
    ws1.cell(row=4, column=1, value="Métrica / Parámetro").fill = header_fill
    ws1.cell(row=4, column=1).font = header_font
    ws1.cell(row=4, column=2, value="Valor").fill = header_fill
    ws1.cell(row=4, column=2).font = header_font

    for idx, (param, val) in enumerate(metrics, start=5):
        ws1.cell(row=idx, column=1, value=param).border = thin_border
        ws1.cell(row=idx, column=2, value=val).border = thin_border

    # Top 5 Procesos
    ws1["A12"] = "2. Top 5 Procesos con Mayor Consumo Promedio"
    ws1["A12"].font = sub_title_font
    
    top5 = df_server.groupby(['Command', 'PID'])['CPU_PERCENT'].agg(['mean', 'max', 'count']).reset_index()
    top5 = top5.sort_values(by='mean', ascending=False).head(5)
    
    headers_top = ["Comando / Proceso", "PID", "%CPU Promedio", "%CPU Máximo", "Muestras Registradas"]
    for c_idx, h in enumerate(headers_top, 1):
        cell = ws1.cell(row=13, column=c_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        
    for r_idx, row in top5.iterrows():
        curr_row = ws1.max_row + 1
        ws1.cell(row=curr_row, column=1, value=row['Command']).border = thin_border
        ws1.cell(row=curr_row, column=2, value=row['PID']).border = thin_border
        ws1.cell(row=curr_row, column=3, value=round(row['mean'], 2)).border = thin_border
        ws1.cell(row=curr_row, column=4, value=round(row['max'], 2)).border = thin_border
        ws1.cell(row=curr_row, column=5, value=row['count']).border = thin_border

    # 2. RESUMEN POR PROCESO
    ws2 = wb.create_sheet(title="Resumen por Proceso")
    ws2.views.sheetView[0].showGridLines = True
    
    grouped = df_server.groupby(['Command', 'PID', 'UID']).agg(
        CPU_Mean=('CPU_PERCENT', 'mean'),
        CPU_Max=('CPU_PERCENT', 'max'),
        CPU_Min=('CPU_PERCENT', 'min'),
        USR_Mean=('USR', 'mean'),
        SYS_Mean=('SYSTEM', 'mean'),
        Samples=('CPU_PERCENT', 'count')
    ).reset_index().sort_values(by='CPU_Mean', ascending=False)
    
    headers_ws2 = ["Comando", "PID", "UID", "%CPU Promedio", "%CPU Máximo", "%CPU Mínimo", "%USR Prom", "%SYS Prom", "Total Muestras"]
    for col, h in enumerate(headers_ws2, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        
    for r_idx, row in grouped.iterrows():
        ws2.append([
            row['Command'], row['PID'], row['UID'],
            round(row['CPU_Mean'], 2), round(row['CPU_Max'], 2), round(row['CPU_Min'], 2),
            round(row['USR_Mean'], 2), round(row['SYS_Mean'], 2), row['Samples']
        ])
    
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, min_col=1, max_col=9):
        for cell in row:
            cell.border = thin_border

    # PIVOT DATOS PARA GRÁFICAS
    top_cmds = top5['Command'].unique()[:5]
    pivot_df = df_server[df_server['Command'].isin(top_cmds)].pivot_table(
        index='Timestamp', columns='Command', values='CPU_PERCENT', aggfunc='max'
    ).fillna(0).reset_index()

    # 3. GRAFICAS POR PROCESO
    ws3 = wb.create_sheet(title="Graficas por Proceso")
    ws3.views.sheetView[0].showGridLines = True
    
    headers_ws3 = list(pivot_df.columns)
    ws3.append(headers_ws3)
    for c in range(1, len(headers_ws3) + 1):
        ws3.cell(row=1, column=c).font = header_font
        ws3.cell(row=1, column=c).fill = header_fill
        
    for r in pivot_df.to_numpy().tolist():
        ws3.append(r)

    chart_row = 2
    for i, cmd_name in enumerate(top_cmds, start=2):
        chart = LineChart()
        chart.title = f"Serie de Tiempo CPU - {cmd_name}"
        chart.style = 13
        chart.y_axis.title = "% CPU"
        chart.x_axis.title = "Tiempo"
        chart.width = 18
        chart.height = 10
        
        data_ref = Reference(ws3, min_col=i, min_row=1, max_row=len(pivot_df)+1)
        cats_ref = Reference(ws3, min_col=1, min_row=2, max_row=len(pivot_df)+1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        ws3.add_chart(chart, f"G{chart_row}")
        chart_row += 18

    # 4. COMPARATIVO TOP PROCESOS (CORREGIDO: Incluye Encabezados de la Tabla)
    ws4 = wb.create_sheet(title="Comparativo Top Procesos")
    ws4.views.sheetView[0].showGridLines = True
    
    headers_ws4 = list(pivot_df.columns)
    ws4.append(headers_ws4)
    
    # Formato visual para los títulos de las columnas
    for c in range(1, len(headers_ws4) + 1):
        cell = ws4.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        
    # Inserción de filas de datos
    for r in pivot_df.to_numpy().tolist():
        ws4.append(r)
        
    comp_chart = LineChart()
    comp_chart.title = f"Comparativo de Consumo de CPU - Top Procesos ({server_name})"
    comp_chart.style = 10
    comp_chart.y_axis.title = "% CPU"
    comp_chart.x_axis.title = "Tiempo"
    comp_chart.width = 24
    comp_chart.height = 14
    
    data_ref_comp = Reference(ws4, min_col=2, max_col=len(top_cmds)+1, min_row=1, max_row=len(pivot_df)+1)
    cats_ref_comp = Reference(ws4, min_col=1, min_row=2, max_row=len(pivot_df)+1)
    comp_chart.add_data(data_ref_comp, titles_from_data=True)
    comp_chart.set_categories(cats_ref_comp)
    ws4.add_chart(comp_chart, "H2")

    # Ajuste automático de ancho de columnas
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(filename)
    print(f" Reporte generado: {filename}")

# --- EJECUCIÓN PRINCIPAL ---
if __name__ == "__main__":
    df_all = parse_all_logs()
    
    if not df_all.empty:
        servers = df_all['Server'].unique()
        print(f"\nSe identificaron {len(servers)} servidor(es): {list(servers)}")
        
        for srv in servers:
            df_srv = df_all[df_all['Server'] == srv]
            generate_excel_for_server(srv, df_srv)
            
        print("\n¡Procesamiento finalizado exitosamente!")